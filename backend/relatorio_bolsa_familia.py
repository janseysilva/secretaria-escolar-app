# -*- coding: utf-8 -*-
"""
Relatorio do Bolsa Familia - cruza fichas de frequencia escolar com a lista
de alunos do Bolsa Familia e calcula a % de presenca de cada um.

Logica portada do programa separado "Relatorio de Presenca"
(gerar_relatorio_presenca_GUI.py), que já é testado e usado pelos colegas do
Jansey - mantida praticamente idêntica de propósito, só sem a parte de
login/atualização automática (que não faz sentido dentro do app de
Secretaria).

Aceita dois modelos de ficha:
  - "Acompanhamento Diario" do sistema Gier (uma ficha por turma, o mes
    inteiro, com o dia marcado presenca/falta/falta justificada);
  - "Relatorio de Frequencia" por materia (uma ficha por turma + componente
    curricular + mes - ex. E.M. Joao Goulart, E.M. Graciele Fernandes),
    somando as materias por aluno para chegar na % geral do mes.

Cruza os PDFs das turmas com uma lista de alunos (.docx ou o formulario do
MEC/Bolsa Familia em .pdf) e gera o relatorio com a % de presenca de cada
aluno da lista, em Excel (.xlsx), Word (.docx) e/ou PDF.

O nome da escola, o mes e o ano letivo sao lidos do proprio PDF, portanto o
programa serve para qualquer unidade de ensino sem alteracao no codigo.
"""

import os
import re
import glob
import difflib
import unicodedata
from datetime import datetime

import pdfplumber
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)


# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

SITUACOES = {"MA", "TR", "TF", "NC"}

# Dia sem aula para o aluno (antes da matricula / apos transferencia).
# Nao conta como dia letivo. Varios tracos Unicode aparecem em PDFs.
TRACOS = {"-", "‐", "‑", "‒", "–", "—", "−"}

RE_DATA = re.compile(r"^\d{2}/\d{2}$")
RE_INICIO_ALUNO = re.compile(r"^\d+\s")
RE_TURMA = re.compile(r"Turma:\s*(.+?)\s+Componente\s+Curricular:", re.IGNORECASE)
RE_TURMA_SIMPLES = re.compile(r"Turma:\s*(\S+)")

# Algumas escolas (ex. CMEI Ismail Aziz) imprimem só uma coluna de total
# ("Faltas", combinando justificadas e normais) e listam as faltas
# justificadas à parte, numa tabela "Faltas Justificadas" no fim do PDF.
RE_TITULO_FALTAS_JUST = re.compile(r"Faltas\s+Justificadas\s*-\s*M[eê]s", re.IGNORECASE)
RE_LINHA_FALTA_JUST = re.compile(r"^\d+\s+(.+?)\s+\d{2}/\d{2}\s+.+$")
RE_ESCOLA = re.compile(r"Escola:\s*(.+?)\s+Ano\s+Letivo:", re.IGNORECASE)
RE_ANO = re.compile(r"Ano\s+Letivo:\s*(\d{4})", re.IGNORECASE)
RE_MES = re.compile(r"M[\wê\xad]*s:\s*([A-Za-zÀ-ú]+)")

# Formulario "Acompanhamento de Frequencia Escolar" do MEC (Bolsa Familia),
# aceito como lista de alunos no lugar do .docx. Um bloco por estudante:
#   [Nome:] FULANO DE TAL      <- com ou sem o prefixo "Nome:"
#   Dt. Nasc.: 27/05/2020
#   NIS: 21344931466
# O nome do aluno e sempre a linha imediatamente ANTES de "Dt. Nasc.:".
# (Algumas escolas imprimem "Nome: FULANO"; outras so "FULANO".)
RE_DT_NASC_MEC = re.compile(r"^\s*Dt\.\s*Nasc", re.IGNORECASE)
RE_PREFIXO_NOME = re.compile(r"^Nome:\s*", re.IGNORECASE)
RE_FORMULARIO_MEC = re.compile(
    r"Acompanhamento\s+de\s+Frequ[êe]ncia\s+Escolar", re.IGNORECASE
)

# Segundo modelo de ficha aceito (ex.: E.M. Joao Goulart, E.M. Graciele
# Fernandes): "RELATÓRIO DE FREQUÊNCIA" por TURMA + COMPONENTE CURRICULAR +
# MES (um arquivo por materia). Frequencia por TEMPO de aula (1T, 2T...), so
# marca falta (F); presenca fica em branco. O % impresso na ficha e de
# FALTAS, nao de presenca. Um aluno tem uma % geral no mes somando as aulas
# e as faltas de TODAS as materias da turma - por isso varios arquivos
# (um por materia) precisam ser lidos e somados antes de virar uma linha.
RE_RELATORIO_FREQUENCIA = re.compile(
    r"RELAT[ÓO]RIO\s+DE\s+FREQU[ÊE]NCIA", re.IGNORECASE
)
TURNOS_MATERIA = ("MATUTINO", "VESPERTINO", "NOTURNO", "INTEGRAL")
MESES_MAIUSC = ("JANEIRO", "FEVEREIRO", "MARÇO", "MARCO", "ABRIL", "MAIO",
                "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO",
                "NOVEMBRO", "DEZEMBRO")

CUTOFF_SIMILARIDADE = 0.87

AZUL = "4472C4"
LARANJA = "F8CBAD"

# Colunas fixas antes dos blocos de cada mes.
COLS_FIXAS_ESQ = [
    ("Nº", "numero", "center"),
    ("Nome do Aluno", "nome", "left"),
    ("Turma", "turma", "center"),
]
# Colunas repetidas dentro do grupo de cada mes.
COLS_MES = [
    ("Dias Letivos", "dias", "center"),
    ("Faltas", "faltas", "center"),
    ("% Presença", "pct", "center"),
]
COL_OBS = ("Observação", "obs", "left")

# Nomes de mes -> numero, para ordenar as colunas em ordem cronologica
# independentemente da ordem em que o usuario adicionou as pastas.
MESES_NUM = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5,
    "junho": 6, "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10,
    "novembro": 11, "dezembro": 12,
}


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def normalizar(nome):
    """Maiusculas, sem acentos, sem pontuacao, espacos colapsados."""
    if not nome:
        return ""
    txt = unicodedata.normalize("NFKD", nome)
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    txt = txt.upper()
    txt = re.sub(r"[^A-Z ]", " ", txt)
    txt = re.sub(r"\s+", " ", txt)
    return txt.strip()


def formatar_percentual(valor):
    """80.0 -> '80,00%'. Vazio continua vazio."""
    if isinstance(valor, (int, float)):
        return "{:.2f}%".format(valor).replace(".", ",")
    return str(valor) if valor else ""


def eh_continuacao_de_nome(tokens):
    """
    Nomes longos quebram em duas linhas no PDF. A linha seguinte e continuacao
    do nome se for so texto em maiusculas, sem token de situacao e sem data.
    """
    if not tokens:
        return False
    for t in tokens:
        if t in SITUACOES:
            return False
        if RE_DATA.match(t):
            return False
        if not re.fullmatch(r"[A-ZÀ-Ú'\.]+", t):
            return False
        # A regua de dias da semana ("S T Q Q S S D S T ...") repetida no
        # cabecalho de cada pagina passaria em todos os testes acima.
        # Sobrenomes tem 2+ letras; letras isoladas nunca sao continuacao.
        if len(t) < 2:
            return False
    return any(len(t) >= 3 for t in tokens)


# ---------------------------------------------------------------------------
# Leitura dos PDFs
# ---------------------------------------------------------------------------

def parsear_linha_aluno(linha, turma):
    """
    Converte uma linha de aluno em dict, ou devolve None se nao for linha de aluno.

    Formato:
      Nº | Estudante | Sit. | Entr. | Saida | 01..30 | Total Falta | Total Falta Comp.
    Ex.:
      19 SAYLON GAMA DOS SANTOS MA 22/12 F F F FJ FJ • F F 8 5
    """
    tokens = linha.split()
    if len(tokens) < 5:
        return None
    if not tokens[0].isdigit():
        return None

    # Localiza a situacao para saber onde o nome termina.
    idx_sit = None
    for i, t in enumerate(tokens):
        if t in SITUACOES:
            idx_sit = i
            break
    if idx_sit is None or idx_sit < 2:
        return None

    nome = " ".join(tokens[1:idx_sit])
    situacao = tokens[idx_sit]

    if not tokens[-1].isdigit():
        return None
    # A maioria das escolas imprime 2 colunas de total no fim da linha
    # (Faltas e Faltas Comp./Justificadas). Algumas (ex. CMEI Ismail Aziz)
    # só têm 1 coluna combinada - nesse caso as justificadas vêm de uma
    # tabela separada "Faltas Justificadas", somadas depois em
    # extrair_dados_pdf.
    if tokens[-2].isdigit():
        n_totais = 2
        total_falta = int(tokens[-2])
        total_falta_comp = int(tokens[-1])
    else:
        n_totais = 1
        total_falta = int(tokens[-1])
        total_falta_comp = 0

    # Depois da situacao vem 1 ou 2 datas DD/MM (entrada e, opcionalmente, saida).
    i = idx_sit + 1
    datas = 0
    while i < len(tokens) - n_totais and datas < 2 and RE_DATA.match(tokens[i]):
        i += 1
        datas += 1

    simbolos = tokens[i:len(tokens) - n_totais]

    # Dias letivos = quantidade de simbolos, ignorando os tracos.
    # Contagem agnostica de glifo: o bullet de presenca pode ser extraido
    # como •, ·, ● etc. dependendo da fonte do PDF.
    dias_letivos = sum(1 for s in simbolos if s not in TRACOS)

    return {
        "numero": int(tokens[0]),
        "nome": nome,
        "situacao": situacao,
        "turma": turma,
        "dias_letivos": dias_letivos,
        "faltas": total_falta,
        "faltas_justificadas": total_falta_comp,
        "simbolos": simbolos,
        # Recontagem dos simbolos, usada apenas para detectar divergencia
        # com os totais impressos (o calculo usa os totais impressos).
        "faltas_contadas": sum(1 for s in simbolos if s == "F"),
        "fj_contadas": sum(1 for s in simbolos if s == "FJ"),
    }


def extrair_cabecalho(texto):
    """Escola, mes e ano letivo, lidos do proprio PDF."""
    escola = RE_ESCOLA.search(texto)
    ano = RE_ANO.search(texto)
    mes = RE_MES.search(texto)
    return {
        "escola": escola.group(1).strip() if escola else "",
        "ano": ano.group(1).strip() if ano else "",
        "mes": mes.group(1).strip().capitalize() if mes else "",
    }


def extrair_dados_pdf(caminho_pdf):
    """Le um PDF de Acompanhamento Diario. Devolve (alunos, info_cabecalho)."""
    alunos = []
    turma_atual = "?"
    info = {"escola": "", "ano": "", "mes": "", "formulario_mec": False}
    faltas_just_por_nome = {}
    em_faltas_justificadas = False

    with pdfplumber.open(caminho_pdf) as pdf:
        for numero, pagina in enumerate(pdf.pages):
            texto = pagina.extract_text() or ""

            # Uma copia do formulario do MEC guardada junto das fichas nao e
            # uma turma: tem 60+ paginas e nenhuma linha de aluno.
            if numero == 0 and RE_FORMULARIO_MEC.search(texto):
                info["formulario_mec"] = True
                return [], info

            if not info["escola"]:
                info.update(extrair_cabecalho(texto))

            for linha in texto.split("\n"):
                linha = linha.strip()
                if not linha:
                    continue

                if RE_TITULO_FALTAS_JUST.search(linha):
                    em_faltas_justificadas = True
                    continue

                # Tabela "Faltas Justificadas" (Nº, Estudante, Data, Motivo):
                # conta quantas faltas justificadas cada aluno teve, para
                # somar depois nas fichas de coluna de total unica.
                if em_faltas_justificadas:
                    m_fj = RE_LINHA_FALTA_JUST.match(linha)
                    if m_fj:
                        nome_fj = m_fj.group(1).strip()
                        faltas_just_por_nome[nome_fj] = (
                            faltas_just_por_nome.get(nome_fj, 0) + 1)
                    continue

                m = RE_TURMA.search(linha) or RE_TURMA_SIMPLES.search(linha)
                if m:
                    turma_atual = " ".join(m.group(1).strip().split())
                    continue

                aluno = parsear_linha_aluno(linha, turma_atual)
                if aluno:
                    alunos.append(aluno)
                    continue

                # Possivel continuacao do nome do aluno anterior.
                if alunos and not RE_INICIO_ALUNO.match(linha):
                    if eh_continuacao_de_nome(linha.split()):
                        alunos[-1]["nome"] += " " + linha.strip()

    # Fichas de coluna de total unica (ex.: CMEI Ismail Aziz): a coluna
    # "Faltas" do quadro principal e o total combinado (justificadas +
    # normais). Usa a tabela "Faltas Justificadas" pra separar os dois
    # numeros, do mesmo jeito que as fichas de duas colunas ja trazem prontos.
    if faltas_just_por_nome:
        for aluno in alunos:
            qtd_just = faltas_just_por_nome.get(aluno["nome"])
            if qtd_just and aluno["faltas_justificadas"] == 0:
                aluno["faltas_justificadas"] = qtd_just
                aluno["faltas"] = max(0, aluno["faltas"] - qtd_just)
                # Os simbolos do quadro principal nao distinguem falta
                # normal de justificada nessas fichas (so marcam "F" pra
                # ambas) - a recontagem de conferencia tem que acompanhar
                # o ajuste, senao acusaria divergencia por engano.
                aluno["faltas_contadas"] = aluno["faltas"]
                aluno["fj_contadas"] = aluno["faltas_justificadas"]

    return alunos, info


def calcular_presenca(aluno):
    """Aplica as regras de calculo e devolve (presencas, percentual, observacao)."""
    dias = aluno["dias_letivos"]
    faltas = aluno["faltas"]
    fj = aluno["faltas_justificadas"]

    # Aluno matriculado no fim do mes: sem simbolo nenhum.
    if dias == 0:
        return None, None, "Sem dias letivos registrados no mes"

    presencas = dias - faltas - fj

    # Inconsistencia do PDF de origem (ocorre em transferidos no meio do mes):
    # o total impresso nao bate com os simbolos da linha.
    if presencas < 0:
        return (
            presencas,
            None,
            "CONFERIR MANUALMENTE: total de faltas impresso no PDF ({} + {}) "
            "maior que os {} dias letivos da linha".format(faltas, fj, dias),
        )

    percentual = round(presencas / dias * 100, 2)

    # Mesma inconsistencia, porem sem estourar para negativo: o percentual sai
    # plausivel e passaria despercebido. Calculo segue pelo total impresso
    # (fonte da verdade), mas a linha e sinalizada para conferencia.
    f_contadas = aluno.get("faltas_contadas")
    fj_contadas = aluno.get("fj_contadas")
    if f_contadas is not None and (f_contadas != faltas or fj_contadas != fj):
        return (
            presencas,
            percentual,
            "CONFERIR MANUALMENTE: o PDF imprime {} falta(s) e {} falta(s) "
            "justificada(s), mas a linha tem {} e {}. Confira na ficha "
            "original".format(faltas, fj, f_contadas, fj_contadas),
        )

    return presencas, percentual, ""


def tipo_de_ficha(caminho_pdf):
    """
    Espia a primeira pagina do PDF para saber que tipo de ficha e:
    'mec'      -> formulario do MEC/Bolsa Familia (e uma lista de alunos,
                  nunca uma turma - pode ter sido salvo por engano na pasta);
    'materia'  -> ficha por materia/tempo de aula (um componente curricular
                  por arquivo - ex. Joao Goulart, Graciele Fernandes);
    'diario'   -> Acompanhamento Diario do Gier (o formato mais comum).
    """
    try:
        with pdfplumber.open(caminho_pdf) as pdf:
            if not pdf.pages:
                return "diario"
            texto = pdf.pages[0].extract_text() or ""
    except Exception:
        return "diario"
    if RE_FORMULARIO_MEC.search(texto):
        return "mec"
    if RE_RELATORIO_FREQUENCIA.search(texto):
        return "materia"
    return "diario"


def extrair_ficha_materia(caminho_pdf):
    """
    Le uma ficha de "Relatorio de Frequencia" por materia (um arquivo por
    turma + componente curricular + mes). Ao contrario do Acompanhamento
    Diario do Gier, aqui a frequencia e por TEMPO de aula (1T, 2T, 4T...),
    so marca falta (F) - presenca fica em branco - e o percentual impresso
    e de FALTAS (faltas / aulas do mes), nao de presenca.

    Devolve (registros, info). `registros` e uma lista de
    {"nome": str, "faltas": int}; as aulas do componente (info["aulas"])
    valem igualmente para todos os alunos da turma nesse componente/mes
    (a ficha nao marca presenca dia a dia, entao nao da para recontar por
    aluno - so o total de faltas impresso e confiavel).
    """
    with pdfplumber.open(caminho_pdf) as pdf:
        linhas = []
        for pagina in pdf.pages:
            linhas.extend((pagina.extract_text() or "").split("\n"))

    info = {"escola": "", "ensino": "", "turno": "", "turma": "",
            "fase": "", "componente": "", "mes": "", "ano": "", "aulas": 0}

    for i, linha in enumerate(linhas):
        if not info["escola"] and linha.strip() == "ESCOLA:" and i + 1 < len(linhas):
            info["escola"] = linhas[i + 1].strip()

        m = re.search(
            r"(Ensino.+?)\s+(" + "|".join(TURNOS_MATERIA) + r")\s+(\S+)\s*$",
            linha,
        )
        if m and not info["turma"]:
            info["ensino"] = m.group(1).strip()
            info["turno"] = m.group(2)
            info["turma"] = m.group(3)

        mc = re.search(
            r"PREFEITURA DE MANAUS\s+(\d+)\s+(.+?)\s+("
            + "|".join(MESES_MAIUSC) + r")\s*$",
            linha, re.IGNORECASE,
        )
        if mc:
            info["fase"] = mc.group(1)
            info["componente"] = mc.group(2).strip()
            info["mes"] = mc.group(3).strip().capitalize().replace("Marco", "Março")

        if linha.strip().startswith("TEMPO"):
            tokens = [t for t in linha.split()[1:] if t not in ("TOTAL", "%")]
            info["aulas"] = len(tokens)

    registros = []
    for linha in linhas:
        if not re.match(r"^\d+\s", linha):
            continue
        tokens = linha.split()
        if len(tokens) < 3:
            continue
        # Os 2 ultimos tokens sao o Total de faltas (inteiro) e o % (x,y).
        if not re.match(r"^\d+(,\d+)?$", tokens[-1]) or not tokens[-2].isdigit():
            continue
        meio = tokens[1:-2]
        while meio and re.fullmatch(r"F[JC]?", meio[-1]):
            meio.pop()
        nome = " ".join(meio)
        if len(nome) < 4:
            continue
        registros.append({"nome": nome, "faltas": int(tokens[-2])})

    return registros, info


def _chave_turma_materia(info):
    """Chave interna para agrupar a mesma turma entre varias fichas de
    materia (varias materias, mesmo mes)."""
    return (normalizar(info.get("ensino", "")), info.get("turno", ""),
            info.get("turma", ""), info.get("fase", ""))


def _rotulo_turma_materia(info):
    """Rotulo legivel da turma, usado na coluna Turma do relatorio."""
    partes = []
    if info.get("fase"):
        partes.append("{}º Ano".format(info["fase"]))
    if info.get("turma"):
        partes.append("Turma {}".format(info["turma"]))
    rotulo = " - ".join(partes) if partes else (info.get("ensino") or "Turma")
    if info.get("turno"):
        rotulo += " ({})".format(info["turno"].capitalize())
    return rotulo


def agregar_fichas_materia(fichas):
    """
    Soma as aulas e as faltas de TODAS as materias de cada aluno, agrupando
    por turma. `fichas` e uma lista de (registros, info) de
    extrair_ficha_materia. Devolve alunos no MESMO formato que
    extrair_dados_pdf produz, para reaproveitar cruzar/calcular_presenca
    sem nenhuma mudanca nelas.
    """
    turmas = {}

    for registros, info in fichas:
        chave_turma = _chave_turma_materia(info)
        turma = turmas.setdefault(chave_turma, {
            "rotulo": _rotulo_turma_materia(info), "alunos": {},
        })
        for registro in registros:
            chave_nome = normalizar(registro["nome"])
            aluno = turma["alunos"].setdefault(chave_nome, {
                "nome": registro["nome"], "aulas": 0, "faltas": 0,
            })
            aluno["aulas"] += info["aulas"]
            aluno["faltas"] += registro["faltas"]

    saida = []
    numero = 0
    for dados_turma in turmas.values():
        for aluno in dados_turma["alunos"].values():
            numero += 1
            saida.append({
                "numero": numero,
                "nome": aluno["nome"],
                "situacao": "MA",
                "turma": dados_turma["rotulo"],
                "dias_letivos": aluno["aulas"],
                "faltas": aluno["faltas"],
                "faltas_justificadas": 0,
            })
    return saida


# ---------------------------------------------------------------------------
# Leitura da lista de alunos (.docx)
# ---------------------------------------------------------------------------

def eh_nome_plausivel(texto):
    """Filtra cabecalhos, numeros soltos e linhas de titulo."""
    if not texto:
        return False
    limpo = texto.strip()
    if len(limpo) < 5:
        return False
    limpo = re.sub(r"^\d+\s*[-.)]?\s*", "", limpo)
    normalizado = normalizar(limpo)
    if len(normalizado.split()) < 2:
        return False
    descartar = {
        "NOME", "NOME DO ALUNO", "ALUNO", "ALUNOS", "TURMA", "BOLSA FAMILIA",
        "LISTA DE ALUNOS", "RELACAO DE ALUNOS", "N NOME",
    }
    if normalizado in descartar:
        return False
    if not re.search(r"[A-Za-zÀ-Ú]{3,}", limpo):
        return False
    return True


def limpar_numeracao(nome):
    return re.sub(r"^\d+\s*[-.)]?\s*", "", nome).strip()


def eh_formulario_mec(caminho_pdf):
    """Reconhece o formulario do MEC pelo titulo da primeira pagina."""
    try:
        with pdfplumber.open(caminho_pdf) as pdf:
            if not pdf.pages:
                return False
            return bool(RE_FORMULARIO_MEC.search(pdf.pages[0].extract_text() or ""))
    except Exception:
        return False


def ler_nomes_pdf_mec(caminho_pdf):
    """
    Le os nomes dos estudantes do formulario do MEC.

    O nome e a linha imediatamente ANTES de "Dt. Nasc.:". Assim funciona tanto
    no formato com prefixo "Nome: FULANO" quanto no formato em que o nome vem
    sozinho na linha (varia de escola para escola).
    """
    linhas = []
    with pdfplumber.open(caminho_pdf) as pdf:
        for pagina in pdf.pages:
            linhas.extend((pagina.extract_text() or "").split("\n"))

    nomes = []
    vistos = set()
    for i, linha in enumerate(linhas):
        if i == 0 or not RE_DT_NASC_MEC.match(linha):
            continue
        nome = RE_PREFIXO_NOME.sub("", linhas[i - 1].strip()).strip()
        if not eh_nome_plausivel(nome):
            continue
        chave = normalizar(nome)
        if chave and chave not in vistos:
            vistos.add(chave)
            nomes.append(nome)

    return nomes


def ler_lista_alunos(caminho, log=print):
    """
    Le a lista de alunos de um .docx ou do formulario do MEC em .pdf.
    O formato e detectado pela extensao.
    """
    extensao = os.path.splitext(caminho)[1].lower()

    if extensao == ".docx":
        log("Lista de alunos: documento do Word")
        nomes = ler_nomes_docx(caminho)
    elif extensao == ".pdf":
        if not eh_formulario_mec(caminho):
            raise RuntimeError(
                "Este PDF não parece o formulário de Acompanhamento de "
                "Frequência Escolar do MEC.\n\nSelecione o formulário do "
                "Bolsa Família ou uma lista em .docx."
            )
        log("Lista de alunos: formulário do MEC (Bolsa Família)")
        nomes = ler_nomes_pdf_mec(caminho)
    else:
        raise RuntimeError(
            "Formato de lista não reconhecido: {}\n\nUse um .docx ou o "
            "formulário do MEC em .pdf.".format(extensao or "sem extensão")
        )

    if not nomes:
        raise RuntimeError(
            "Nenhum nome de aluno foi encontrado em:\n{}".format(
                os.path.basename(caminho)
            )
        )
    return nomes


def ler_nomes_docx(caminho_docx):
    """Le nomes dos paragrafos e tambem das tabelas do .docx."""
    documento = docx.Document(caminho_docx)
    brutos = []

    for paragrafo in documento.paragraphs:
        brutos.append(paragrafo.text)

    for tabela in documento.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                brutos.append(celula.text)

    nomes = []
    vistos = set()
    for texto in brutos:
        for parte in texto.split("\n"):
            nome = parte.strip()
            if not eh_nome_plausivel(nome):
                continue
            chave = normalizar(nome)
            if chave and chave not in vistos:
                vistos.add(chave)
                nomes.append(nome.strip())

    return nomes


# ---------------------------------------------------------------------------
# Cruzamento lista x PDFs
# ---------------------------------------------------------------------------

def _celula_vazia():
    return {
        "dias_letivos": "-", "faltas": "-", "faltas_justificadas": "-",
        "presencas": "-", "percentual": "-", "encontrado": False,
    }


def cruzar(nomes_lista, meses_alunos, rotulos):
    """
    Cruza os nomes da lista com os alunos extraidos dos PDFs de cada mes.

    `meses_alunos` e uma lista paralela a `rotulos`: cada item e a lista de
    alunos daquele mes. Para cada nome da lista, busca em cada mes por
    correspondencia exata e, depois, aproximada (difflib, cutoff 0.87).

    Devolve a lista de resultados; cada um tem o campo "meses" com um dicionario
    por mes (na mesma ordem de `rotulos`).
    """
    n = len(meses_alunos)
    prefixar = n > 1

    indices = []
    listas_chaves = []
    for alunos in meses_alunos:
        indice = {}
        for aluno in alunos:
            chave = normalizar(aluno["nome"])
            if chave and chave not in indice:
                indice[chave] = aluno
        indices.append(indice)
        listas_chaves.append(list(indice.keys()))

    resultados = []
    for nome_lista in nomes_lista:
        nome_limpo = limpar_numeracao(nome_lista)
        chave = normalizar(nome_limpo)

        turma = "-"
        meses_data = []
        notas_por_mes = []
        encontrado_algum = False

        for i in range(n):
            aluno = indices[i].get(chave)
            aprox = ""
            if aluno is None:
                proximos = difflib.get_close_matches(
                    chave, listas_chaves[i], n=1, cutoff=CUTOFF_SIMILARIDADE
                )
                if proximos:
                    aluno = indices[i][proximos[0]]
                    aprox = "OK (correspondência aproximada: {})".format(
                        aluno["nome"]
                    )

            if aluno is None:
                meses_data.append(_celula_vazia())
                notas_por_mes.append([])
                continue

            encontrado_algum = True
            if turma == "-":
                turma = aluno["turma"]

            presencas, percentual, observacao = calcular_presenca(aluno)
            meses_data.append({
                "dias_letivos": aluno["dias_letivos"],
                "faltas": aluno["faltas"],
                "faltas_justificadas": aluno["faltas_justificadas"],
                "presencas": presencas if presencas is not None else "-",
                "percentual": percentual if percentual is not None else "",
                "encontrado": True,
            })

            notas = []
            if aprox:
                notas.append(aprox)
            if observacao:
                notas.append(observacao)
            if aluno["situacao"] == "TR":
                notas.append("Aluno transferido")
            elif aluno["situacao"] == "TF":
                notas.append("Transferido para fora da rede")
            elif aluno["situacao"] == "NC":
                notas.append("Situação: não comparecimento")

            prefixo = (rotulos[i] + ": ") if prefixar else ""
            notas_por_mes.append([prefixo + nota for nota in notas])

        if not encontrado_algum:
            observacao_final = "NÃO LOCALIZADO na base de PDFs"
        else:
            partes = []
            for i in range(n):
                if prefixar and not meses_data[i]["encontrado"]:
                    partes.append("Não localizado em " + rotulos[i])
                partes.extend(notas_por_mes[i])
            observacao_final = " | ".join(partes)

        resultados.append({
            "nome": nome_limpo,
            "turma": turma,
            "observacao": observacao_final,
            "meses": meses_data,
        })

    # Ordem alfabetica por nome, com os nao localizados na posicao correta.
    resultados.sort(key=lambda r: normalizar(r["nome"]))
    return resultados


def abaixo_no_mes(resultado, i, percentual_minimo):
    """True se o mes i tem % de presenca abaixo do minimo."""
    percentual = resultado["meses"][i]["percentual"]
    return (
        isinstance(percentual, (int, float)) and percentual < percentual_minimo
    )


def resumir(resultados, percentual_minimo, n_meses):
    """Numeros usados no resumo dos relatorios e no painel final da janela."""
    abaixo_por_mes = [
        [r for r in resultados if abaixo_no_mes(r, i, percentual_minimo)]
        for i in range(n_meses)
    ]
    abaixo = [
        r for r in resultados
        if any(abaixo_no_mes(r, i, percentual_minimo) for i in range(n_meses))
    ]
    return {
        "total": len(resultados),
        "abaixo": abaixo,
        "abaixo_por_mes": abaixo_por_mes,
        "aproximados": [r for r in resultados if "aproximada" in r["observacao"]],
        "conferir": [r for r in resultados if "CONFERIR MANUALMENTE" in r["observacao"]],
        "nao_localizados": [
            r for r in resultados if "NÃO LOCALIZADO na base" in r["observacao"]
        ],
    }


# -- modelo de colunas (dinamico conforme o numero de meses) ----------------

def colunas_folha(n_meses):
    """Lista das colunas finais (folhas) da tabela, na ordem de exibicao."""
    cols = list(COLS_FIXAS_ESQ)
    for i in range(n_meses):
        for titulo, tipo, align in COLS_MES:
            cols.append((titulo, (tipo, i), align))
    cols.append(COL_OBS)
    return cols


def indices_colunas_mes(i):
    """Indices (0-based) das 3 colunas do mes i entre as colunas folha."""
    inicio = len(COLS_FIXAS_ESQ) + i * len(COLS_MES)
    return list(range(inicio, inicio + len(COLS_MES)))


def valor_coluna(chave, numero, resultado, para_texto=False):
    """Valor de uma celula, dada a chave da coluna."""
    if chave == "numero":
        return numero
    if chave == "nome":
        return resultado["nome"]
    if chave == "turma":
        return resultado["turma"]
    if chave == "obs":
        return resultado["observacao"]
    tipo, i = chave
    mes = resultado["meses"][i]
    if tipo == "dias":
        return mes["dias_letivos"]
    if tipo == "faltas":
        return mes["faltas"]
    # tipo == "pct"
    if para_texto:
        return formatar_percentual(mes["percentual"])
    return mes["percentual"]


def rotulo_periodo(meses):
    """Ex.: 'Junho e Julho / 2026' ou 'Junho / 2026'."""
    nomes = [m["rotulo"] for m in meses if m.get("rotulo")]
    anos = sorted({m["ano"] for m in meses if m.get("ano")})
    if not nomes:
        base = "Frequência Mensal"
    elif len(nomes) == 1:
        base = nomes[0]
    else:
        base = ", ".join(nomes[:-1]) + " e " + nomes[-1]
    if anos:
        base += " / " + "-".join(anos)
    return base


def titulo_relatorio(meses):
    return "Relatório de Frequência - " + rotulo_periodo(meses)


def linha_resumo(resumo, meses, percentual_minimo):
    """Texto do resumo (emitido em / totais) para o cabecalho dos relatorios."""
    partes = ["{} aluno(s) na lista".format(resumo["total"])]
    if len(meses) <= 1:
        partes.append("{} abaixo de {:g}% de presença".format(
            len(resumo["abaixo_por_mes"][0]) if resumo["abaixo_por_mes"] else 0,
            percentual_minimo,
        ))
    else:
        detalhe = ", ".join(
            "{}: {}".format(meses[i]["rotulo"], len(resumo["abaixo_por_mes"][i]))
            for i in range(len(meses))
        )
        partes.append("abaixo de {:g}% — {}".format(percentual_minimo, detalhe))
    partes.append("{} não localizado(s)".format(len(resumo["nao_localizados"])))
    return "  |  ".join(partes)


# ---------------------------------------------------------------------------
# Saida: Excel
# ---------------------------------------------------------------------------

def gerar_xlsx(resultados, caminho_saida, percentual_minimo, meses):
    wb = Workbook()
    ws = wb.active
    ws.title = "Presença"

    n = len(meses)
    colunas = colunas_folha(n)
    ncols = len(colunas)
    base = len(COLS_FIXAS_ESQ)

    fonte_cabecalho = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    fundo_cabecalho = PatternFill("solid", fgColor=AZUL)
    fundo_alerta = PatternFill("solid", fgColor=LARANJA)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Larguras: Nº, Nome, Turma, [Dias, Faltas, %]xN, Observação
    larguras = [6, 38, 12] + [11, 8, 12] * n + [42]
    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura

    def estilo_cabecalho(celula, texto):
        celula.value = texto
        celula.font = fonte_cabecalho
        celula.fill = fundo_cabecalho
        celula.alignment = centro

    # Linha 1 = grupos; linha 2 = subtitulos. Colunas fixas ocupam as duas.
    for c in range(len(COLS_FIXAS_ESQ)):
        ws.merge_cells(start_row=1, start_column=c + 1, end_row=2, end_column=c + 1)
        estilo_cabecalho(ws.cell(row=1, column=c + 1), COLS_FIXAS_ESQ[c][0])
        ws.cell(row=2, column=c + 1).fill = fundo_cabecalho

    for i, mes in enumerate(meses):
        c0 = base + i * len(COLS_MES) + 1
        ws.merge_cells(start_row=1, start_column=c0, end_row=1,
                       end_column=c0 + len(COLS_MES) - 1)
        estilo_cabecalho(ws.cell(row=1, column=c0), mes["rotulo"])
        for k in range(len(COLS_MES)):
            ws.cell(row=1, column=c0 + k).fill = fundo_cabecalho
            estilo_cabecalho(ws.cell(row=2, column=c0 + k), COLS_MES[k][0])

    ws.merge_cells(start_row=1, start_column=ncols, end_row=2, end_column=ncols)
    estilo_cabecalho(ws.cell(row=1, column=ncols), COL_OBS[0])
    ws.cell(row=2, column=ncols).fill = fundo_cabecalho

    # Dados (comecam na linha 3)
    for numero, resultado in enumerate(resultados, start=1):
        linha = numero + 2
        for cidx, (titulo, chave, align) in enumerate(colunas, start=1):
            celula = ws.cell(row=linha, column=cidx,
                             value=valor_coluna(chave, numero, resultado))
            celula.font = Font(name="Arial", size=11)
            celula.alignment = Alignment(horizontal=align, vertical="center")

        # Destaca o bloco (Dias/Faltas/%) do mes que estiver abaixo do minimo.
        for i in range(n):
            if abaixo_no_mes(resultado, i, percentual_minimo):
                for leaf in indices_colunas_mes(i):
                    ws.cell(row=linha, column=leaf + 1).fill = fundo_alerta

    ws.freeze_panes = "A3"
    wb.save(caminho_saida)


# ---------------------------------------------------------------------------
# Saida: Word
# ---------------------------------------------------------------------------

def _sombrear_celula(celula, cor_hex):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:fill"), cor_hex)
    celula._tc.get_or_add_tcPr().append(sombra)


def _texto_celula(celula, texto, negrito=False, branco=False, tamanho=7,
                  centralizado=False):
    celula.text = ""
    paragrafo = celula.paragraphs[0]
    paragrafo.paragraph_format.space_before = Pt(1)
    paragrafo.paragraph_format.space_after = Pt(1)
    if centralizado:
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trecho = paragrafo.add_run(str(texto))
    trecho.font.size = Pt(tamanho)
    trecho.font.name = "Arial"
    trecho.font.bold = negrito
    if branco:
        trecho.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _marcar_cabecalho_repetido(linha):
    """Faz a linha se repetir como cabecalho no topo de cada pagina impressa."""
    props = linha._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    props.append(repetir)


def _larguras_docx(n_meses):
    """Larguras (Cm) das colunas folha. Soma ~27,3 cm (A4 paisagem c/ margens)."""
    larguras = [Cm(1.1), Cm(5.2), Cm(2.7)]
    usado = 1.1 + 5.2 + 2.7
    for _ in range(n_meses):
        larguras += [Cm(1.5), Cm(1.2), Cm(1.7)]
        usado += 1.5 + 1.2 + 1.7
    larguras.append(Cm(max(27.3 - usado, 3.0)))
    return larguras


def _forcar_layout_fixo(tabela, larguras):
    """
    Fixa a grade de colunas da tabela do Word. Sem isso o Word recalcula as
    larguras pelo conteudo e a tabela transborda a pagina.
    """
    tbl = tabela._tbl
    tblPr = tbl.tblPr

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    largura_total = OxmlElement("w:tblW")
    largura_total.set(qn("w:w"), str(int(sum(l.twips for l in larguras))))
    largura_total.set(qn("w:type"), "dxa")
    tblPr.append(largura_total)

    grid = tbl.find(qn("w:tblGrid"))
    for coluna in list(grid.findall(qn("w:gridCol"))):
        grid.remove(coluna)
    for largura in larguras:
        coluna = OxmlElement("w:gridCol")
        coluna.set(qn("w:w"), str(int(largura.twips)))
        grid.append(coluna)


def gerar_docx(resultados, caminho_saida, percentual_minimo, meses):
    documento = docx.Document()

    secao = documento.sections[0]
    secao.orientation = WD_ORIENT.LANDSCAPE
    secao.page_width, secao.page_height = secao.page_height, secao.page_width
    for margem in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(secao, margem, Cm(1.2))

    n = len(meses)
    colunas = colunas_folha(n)
    ncols = len(colunas)
    base = len(COLS_FIXAS_ESQ)
    larguras = _larguras_docx(n)

    escola = meses[0].get("escola") if meses else ""
    if escola:
        p = documento.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t = p.add_run(escola)
        t.font.size = Pt(12)
        t.font.bold = True
        t.font.name = "Arial"

    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = p.add_run(titulo_relatorio(meses))
    t.font.size = Pt(14)
    t.font.bold = True
    t.font.name = "Arial"
    t.font.color.rgb = RGBColor(0x2F, 0x52, 0x8F)

    resumo = resumir(resultados, percentual_minimo, n)
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = p.add_run("Emitido em {}  |  {}".format(
        datetime.now().strftime("%d/%m/%Y às %H:%M"),
        linha_resumo(resumo, meses, percentual_minimo),
    ))
    t.font.size = Pt(8)
    t.font.name = "Arial"
    t.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    documento.add_paragraph()

    # Duas linhas de cabecalho: grupos (mes) e subtitulos.
    tabela = documento.add_table(rows=2, cols=ncols)
    tabela.style = "Table Grid"
    tabela.autofit = False
    tabela.allow_autofit = False

    def celula_cab(celula, texto):
        _sombrear_celula(celula, AZUL)
        _texto_celula(celula, texto, negrito=True, branco=True, centralizado=True)

    # Colunas fixas: mescla vertical das duas linhas.
    for c in range(base):
        cel = tabela.cell(0, c).merge(tabela.cell(1, c))
        celula_cab(cel, COLS_FIXAS_ESQ[c][0])
    # Grupos de mes: mescla horizontal na linha 0; subtitulos na linha 1.
    for i, mes in enumerate(meses):
        c0 = base + i * len(COLS_MES)
        grupo = tabela.cell(0, c0).merge(tabela.cell(0, c0 + len(COLS_MES) - 1))
        celula_cab(grupo, mes["rotulo"])
        for k in range(len(COLS_MES)):
            celula_cab(tabela.cell(1, c0 + k), COLS_MES[k][0])
    # Observação: mescla vertical.
    cel_obs = tabela.cell(0, ncols - 1).merge(tabela.cell(1, ncols - 1))
    celula_cab(cel_obs, COL_OBS[0])

    _forcar_layout_fixo(tabela, larguras)
    _marcar_cabecalho_repetido(tabela.rows[0])
    _marcar_cabecalho_repetido(tabela.rows[1])

    for numero, resultado in enumerate(resultados, start=1):
        celulas = tabela.add_row().cells
        destaque_mes = {
            leaf for i in range(n) if abaixo_no_mes(resultado, i, percentual_minimo)
            for leaf in indices_colunas_mes(i)
        }
        for cidx, (titulo, chave, align) in enumerate(colunas):
            valor = valor_coluna(chave, numero, resultado, para_texto=True)
            _texto_celula(celulas[cidx], valor, centralizado=(align == "center"))
            celulas[cidx].width = larguras[cidx]
            if cidx in destaque_mes:
                _sombrear_celula(celulas[cidx], LARANJA)

    documento.add_paragraph()
    p = documento.add_paragraph()
    t = p.add_run(
        "Células destacadas em laranja: mês com presença abaixo de {:g}%."
        .format(percentual_minimo)
    )
    t.font.size = Pt(8)
    t.font.name = "Arial"

    documento.add_paragraph()
    documento.add_paragraph()
    p = documento.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t = p.add_run("_" * 55 + "\nResponsável pelas informações")
    t.font.size = Pt(9)
    t.font.name = "Arial"

    documento.save(caminho_saida)


# ---------------------------------------------------------------------------
# Saida: PDF
# ---------------------------------------------------------------------------

def _rodape_pdf(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.grey)
    canvas.drawRightString(
        landscape(A4)[0] - 12 * mm, 8 * mm, "Página {}".format(doc.page)
    )
    canvas.restoreState()


def _larguras_pdf(n_meses):
    """Larguras (pt) das colunas folha; a Observação fica com o que sobrar."""
    usavel = landscape(A4)[0] - 20 * mm
    larguras = [24, 150, 58]
    for _ in range(n_meses):
        larguras += [40, 34, 46]
    larguras.append(max(usavel - sum(larguras), 90))
    return larguras


def gerar_pdf(resultados, caminho_saida, percentual_minimo, meses):
    documento = SimpleDocTemplate(
        caminho_saida,
        pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm,
        topMargin=10 * mm, bottomMargin=14 * mm,
        title=titulo_relatorio(meses),
    )

    estilos = getSampleStyleSheet()
    estilo_escola = ParagraphStyle(
        "escola", parent=estilos["Normal"], fontName="Helvetica-Bold",
        fontSize=11, alignment=1, spaceAfter=2,
    )
    estilo_titulo = ParagraphStyle(
        "titulo", parent=estilos["Normal"], fontName="Helvetica-Bold",
        fontSize=14, alignment=1, textColor=colors.HexColor("#2F528F"),
        spaceAfter=3,
    )
    estilo_resumo = ParagraphStyle(
        "resumo", parent=estilos["Normal"], fontName="Helvetica",
        fontSize=7.5, alignment=1, textColor=colors.grey, spaceAfter=8,
    )
    estilo_celula = ParagraphStyle(
        "celula", parent=estilos["Normal"], fontName="Helvetica",
        fontSize=6.5, leading=8,
    )
    estilo_celula_centro = ParagraphStyle(
        "celula_c", parent=estilo_celula, alignment=1,
    )
    estilo_cabecalho = ParagraphStyle(
        "cab", parent=estilos["Normal"], fontName="Helvetica-Bold",
        fontSize=7, leading=8.5, alignment=1, textColor=colors.white,
    )
    estilo_nota = ParagraphStyle(
        "nota", parent=estilos["Normal"], fontName="Helvetica",
        fontSize=7, spaceBefore=8,
    )

    n = len(meses)
    colunas = colunas_folha(n)
    ncols = len(colunas)
    base = len(COLS_FIXAS_ESQ)

    elementos = []
    escola = meses[0].get("escola") if meses else ""
    if escola:
        elementos.append(Paragraph(escola, estilo_escola))
    elementos.append(Paragraph(titulo_relatorio(meses), estilo_titulo))

    resumo = resumir(resultados, percentual_minimo, n)
    elementos.append(Paragraph(
        "Emitido em {} &nbsp;|&nbsp; {}".format(
            datetime.now().strftime("%d/%m/%Y às %H:%M"),
            linha_resumo(resumo, meses, percentual_minimo).replace(
                "|", "&nbsp;|&nbsp;"),
        ),
        estilo_resumo,
    ))

    # Duas linhas de cabecalho: grupos (mes) e subtitulos, com SPAN.
    grupo = [""] * ncols
    sub = [""] * ncols
    comandos_span = []
    for c in range(base):
        grupo[c] = COLS_FIXAS_ESQ[c][0]
        comandos_span.append(("SPAN", (c, 0), (c, 1)))
    for i, mes in enumerate(meses):
        c0 = base + i * len(COLS_MES)
        grupo[c0] = mes["rotulo"]
        comandos_span.append(("SPAN", (c0, 0), (c0 + len(COLS_MES) - 1, 0)))
        for k in range(len(COLS_MES)):
            sub[c0 + k] = COLS_MES[k][0]
    grupo[ncols - 1] = COL_OBS[0]
    comandos_span.append(("SPAN", (ncols - 1, 0), (ncols - 1, 1)))

    dados = [
        [Paragraph(x, estilo_cabecalho) for x in grupo],
        [Paragraph(x, estilo_cabecalho) for x in sub],
    ]

    celulas_destacadas = []
    for numero, resultado in enumerate(resultados, start=1):
        linha = []
        for titulo, chave, align in colunas:
            valor = valor_coluna(chave, numero, resultado, para_texto=True)
            estilo_cel = estilo_celula_centro if align == "center" else estilo_celula
            linha.append(Paragraph(str(valor), estilo_cel))
        dados.append(linha)
        r = len(dados) - 1
        for i in range(n):
            if abaixo_no_mes(resultado, i, percentual_minimo):
                for leaf in indices_colunas_mes(i):
                    celulas_destacadas.append((leaf, r))

    tabela = Table(dados, colWidths=_larguras_pdf(n), repeatRows=2)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#" + AZUL)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#F2F5FB")]),
    ] + comandos_span
    for coluna, linha_i in celulas_destacadas:
        estilo.append(("BACKGROUND", (coluna, linha_i), (coluna, linha_i),
                       colors.HexColor("#" + LARANJA)))
    tabela.setStyle(TableStyle(estilo))

    elementos.append(tabela)
    elementos.append(Paragraph(
        "Células destacadas em laranja: mês com presença abaixo de {:g}%."
        .format(percentual_minimo),
        estilo_nota,
    ))
    elementos.append(Spacer(1, 30))

    # Linha continua desenhada. Uma sequencia de "_" sai falhada no PDF.
    assinatura = Table(
        [[""], ["Responsável pelas informações"]], colWidths=[230],
        hAlign="CENTER",
    )
    assinatura.setStyle(TableStyle([
        ("LINEABOVE", (0, 1), (0, 1), 0.7, colors.black),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 1), (0, 1), "Helvetica"),
        ("FONTSIZE", (0, 1), (0, 1), 8.5),
        ("TOPPADDING", (0, 1), (0, 1), 5),
    ]))
    elementos.append(assinatura)

    documento.build(elementos, onFirstPage=_rodape_pdf, onLaterPages=_rodape_pdf)


# ---------------------------------------------------------------------------
# Processamento completo
# ---------------------------------------------------------------------------

GERADORES = {
    "xlsx": ("Excel", gerar_xlsx),
    "docx": ("Word", gerar_docx),
    "pdf": ("PDF", gerar_pdf),
}


def sugerir_nome(meses):
    """Nome de arquivo sugerido, montado com os meses e o ano lidos dos PDFs."""
    partes = ["Presenca"]
    partes += [m["rotulo"] for m in meses if m.get("rotulo")]
    partes += sorted({m["ano"] for m in meses if m.get("ano")})
    nome = "_".join(partes) if len(partes) > 1 else "Presenca"
    # Tira acentos e o que o Windows nao aceita em nome de arquivo.
    nome = unicodedata.normalize("NFKD", nome)
    nome = "".join(c for c in nome if not unicodedata.combining(c))
    return re.sub(r'[<>:"/\\|?*\s]', "_", nome)


def detectar_mes_pasta(pasta):
    """
    Espia o primeiro PDF de turma da pasta e devolve o mes do cabecalho
    (ex.: 'Junho'), para mostrar na janela ao adicionar a pasta. '' se nao achar.
    """
    for caminho in sorted(glob.glob(os.path.join(pasta, "*.pdf"))):
        try:
            tipo = tipo_de_ficha(caminho)
            if tipo == "mec":
                continue  # e a lista do MEC, nao uma turma
            if tipo == "materia":
                _, info = extrair_ficha_materia(caminho)
            else:
                with pdfplumber.open(caminho) as pdf:
                    if not pdf.pages:
                        continue
                    texto = pdf.pages[0].extract_text() or ""
                info = extrair_cabecalho(texto)
            if info.get("mes"):
                return info["mes"]
        except Exception:
            continue
    return ""


def mes_do_nome(texto):
    """Extrai um nome de mes reconhecido de um texto (ex.: nome da pasta)."""
    normalizado = normalizar(texto).lower()
    for nome in MESES_NUM:
        if re.search(r"\b" + nome + r"\b", normalizado):
            return nome
    return ""


def _chave_ordem_mes(mes, posicao):
    """Ordena os meses por ano e numero do mes; desconhecidos vao ao fim."""
    ano = int(mes["ano"]) if mes.get("ano", "").isdigit() else 9999
    num = MESES_NUM.get(mes_do_nome(mes.get("rotulo", "")), 99)
    return (ano, num, posicao)


def ler_pasta_mes(pasta, lista_abs, log=print):
    """
    Le todas as fichas de turma de uma pasta (um mes). Aceita tanto o
    Acompanhamento Diario do Gier (uma ficha por turma, com o mes inteiro)
    quanto fichas por materia/tempo de aula (uma ficha por turma + componente
    curricular + mes - ex. Joao Goulart, Graciele Fernandes), somando as
    materias por aluno antes de entrar no relatorio.
    Devolve (alunos, info).
    """
    pdfs = sorted(glob.glob(os.path.join(pasta, "*.pdf")))
    if not pdfs:
        raise RuntimeError("Nenhum PDF encontrado na pasta:\n{}".format(pasta))

    # A lista de alunos costuma ficar guardada junto dos PDFs. Se for o
    # formulario do MEC (que tambem e .pdf), nao pode entrar como turma.
    pdfs = [p for p in pdfs if os.path.normcase(os.path.abspath(p)) != lista_abs]
    if not pdfs:
        raise RuntimeError(
            "A pasta só contém o arquivo da lista de alunos, nenhuma ficha "
            "de turma:\n{}".format(pasta)
        )

    alunos = []
    fichas_materia = []
    info = {"escola": "", "ano": "", "mes": ""}

    for caminho in pdfs:
        nome_arquivo = os.path.basename(caminho)
        try:
            tipo = tipo_de_ficha(caminho)

            if tipo == "mec":
                log("   {} -> ignorado (é o formulário do MEC, não uma turma)"
                    .format(nome_arquivo))

            elif tipo == "materia":
                registros, info_pdf = extrair_ficha_materia(caminho)
                fichas_materia.append((registros, info_pdf))
                if not info["mes"] and info_pdf.get("mes"):
                    info["mes"] = info_pdf["mes"]
                    info["ano"] = info_pdf.get("ano", "")
                if not info["escola"] and info_pdf.get("escola"):
                    info["escola"] = info_pdf["escola"]
                log("   {} -> {} aluno(s) em {} ({})".format(
                    nome_arquivo, len(registros),
                    info_pdf.get("componente") or "?", info_pdf.get("mes") or "?"))

            else:
                alunos_pdf, info_pdf = extrair_dados_pdf(caminho)
                if info_pdf.get("formulario_mec"):
                    log("   {} -> ignorado (é o formulário do MEC, não uma turma)"
                        .format(nome_arquivo))
                else:
                    alunos.extend(alunos_pdf)
                    if not info["mes"] and info_pdf.get("mes"):
                        info["mes"] = info_pdf["mes"]
                        info["ano"] = info_pdf.get("ano", "")
                    if not info["escola"] and info_pdf.get("escola"):
                        info["escola"] = info_pdf["escola"]
                    log("   {} -> {} aluno(s)".format(nome_arquivo, len(alunos_pdf)))
        except Exception as excecao:
            # Um PDF ilegivel (aberto em outro programa, sincronizando na nuvem
            # ou corrompido) nao pode derrubar o relatorio inteiro.
            log("   {} -> IGNORADO: não foi possível ler ({})".format(
                nome_arquivo, excecao))

    if fichas_materia:
        materias = sorted({
            info_pdf.get("componente") or "?" for _, info_pdf in fichas_materia
        })
        log("   Somando {} componente(s) por aluno: {}".format(
            len(materias), ", ".join(materias)))
        alunos.extend(agregar_fichas_materia(fichas_materia))

    if not alunos:
        raise RuntimeError(
            "Nenhuma ficha de frequência foi encontrada na pasta:\n"
            "{}\n\nConfira se selecionou a pasta certa.".format(pasta)
        )
    return alunos, info


def analisar(pastas, arquivo_lista, percentual_minimo, log=print):
    """
    Le os PDFs de cada pasta (um mes por pasta) e cruza com a lista, sem gravar.
    Devolve (resultados, meses, resumo).
    """
    if isinstance(pastas, str):
        pastas = [pastas]
    if not pastas:
        raise RuntimeError("Selecione pelo menos a pasta de um mês.")

    lista_abs = os.path.normcase(os.path.abspath(arquivo_lista))
    meses = []
    meses_alunos = []

    for indice, pasta in enumerate(pastas):
        log("Lendo os PDFs de: {}".format(os.path.basename(pasta.rstrip("\\/")) or pasta))
        alunos, info = ler_pasta_mes(pasta, lista_abs, log)

        rotulo = info.get("mes") or os.path.basename(pasta.rstrip("\\/")) \
            or "Mês {}".format(indice + 1)

        # O nome da pasta discorda do mes impresso nos PDFs? Provavel exportacao
        # do mes errado no Gier (foi assim que descobrimos um "Julho" que na
        # verdade trazia os dados de Junho).
        mes_pasta = mes_do_nome(os.path.basename(pasta.rstrip("\\/")))
        mes_pdf = mes_do_nome(info.get("mes", ""))
        if mes_pasta and mes_pdf and mes_pasta != mes_pdf:
            log("   ATENÇÃO: a pasta parece ser de {} mas os PDFs dizem "
                "'Mês: {}'. Confira se exportou o mês certo no Gier."
                .format(mes_pasta.capitalize(), info.get("mes")))

        meses.append({
            "rotulo": rotulo, "ano": info.get("ano", ""),
            "escola": info.get("escola", ""), "pasta": pasta,
        })
        meses_alunos.append(alunos)
        log("   => {}: {} aluno(s) no total".format(rotulo, len(alunos)))

    # Duas pastas com o mesmo mes? Quase sempre e o mesmo mes exportado duas
    # vezes. Avisa e desambigua os rotulos para as colunas nao colidirem.
    contagem = {}
    for mes in meses:
        contagem[mes["rotulo"]] = contagem.get(mes["rotulo"], 0) + 1
    if any(qtd > 1 for qtd in contagem.values()):
        repetidos = ", ".join(r for r, q in contagem.items() if q > 1)
        log("\n*** ATENÇÃO: mais de uma pasta tem o mesmo mês ({}). Verifique se "
            "você não exportou o mesmo mês duas vezes no Gier. ***".format(repetidos))
        vistos = {}
        for mes in meses:
            vistos[mes["rotulo"]] = vistos.get(mes["rotulo"], 0) + 1
            if vistos[mes["rotulo"]] > 1:
                mes["rotulo"] = "{} ({})".format(mes["rotulo"], vistos[mes["rotulo"]])

    # Ordena as colunas por mes (Junho antes de Julho), qualquer que tenha
    # sido a ordem em que o usuario adicionou as pastas.
    ordem = sorted(range(len(meses)),
                   key=lambda k: _chave_ordem_mes(meses[k], k))
    meses = [meses[k] for k in ordem]
    meses_alunos = [meses_alunos[k] for k in ordem]

    log("")
    if meses and meses[0].get("escola"):
        log("Escola: {}".format(meses[0]["escola"]))
    log("Meses: {}".format(", ".join(m["rotulo"] for m in meses)))

    log("\nLendo a lista de alunos...")
    nomes_lista = ler_lista_alunos(arquivo_lista, log=log)
    log("Nomes na lista: {}".format(len(nomes_lista)))

    log("\nCruzando os dados...")
    rotulos = [m["rotulo"] for m in meses]
    resultados = cruzar(nomes_lista, meses_alunos, rotulos)
    resumo = resumir(resultados, percentual_minimo, len(meses))

    if resumo["aproximados"]:
        log("\nCorrespondências aproximadas ({}) - confira se é a mesma criança:"
            .format(len(resumo["aproximados"])))
        for r in resumo["aproximados"]:
            log("   {}".format(r["nome"]))

    if resumo["nao_localizados"]:
        log("\nNão localizados em nenhum mês ({}):".format(
            len(resumo["nao_localizados"])))
        for r in resumo["nao_localizados"]:
            log("   {}".format(r["nome"]))

    if resumo["conferir"]:
        log("\nConferir manualmente na ficha original ({}):".format(
            len(resumo["conferir"])))
        for r in resumo["conferir"]:
            log("   {} ({})".format(r["nome"], r["turma"]))

    for i, mes in enumerate(meses):
        log("Abaixo de {:g}% em {}: {} aluno(s)".format(
            percentual_minimo, mes["rotulo"], len(resumo["abaixo_por_mes"][i])))

    return resultados, meses, resumo


def salvar(resultados, caminho_base, percentual_minimo, meses, formatos, log=print):
    """Grava um arquivo por formato pedido. `caminho_base` e sem extensao."""
    if not formatos:
        raise RuntimeError("Escolha pelo menos um formato de saída.")

    log("\nGerando arquivo(s)...")
    arquivos = []
    for formato in formatos:
        rotulo, gerador = GERADORES[formato]
        caminho = "{}.{}".format(caminho_base, formato)
        gerador(resultados, caminho, percentual_minimo, meses)
        arquivos.append(caminho)
        log("   {} salvo: {}".format(rotulo, os.path.basename(caminho)))

    return arquivos


def processar(pastas, arquivo_lista, caminho_base, percentual_minimo,
              formatos=("xlsx",), log=print):
    """Analisa e grava de uma vez."""
    resultados, meses, resumo = analisar(pastas, arquivo_lista, percentual_minimo, log=log)
    arquivos = salvar(resultados, caminho_base, percentual_minimo, meses, formatos, log=log)
    return resultados, arquivos, resumo
